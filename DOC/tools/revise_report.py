"""Revise the existing Word report, retaining its styles, sections and diagrams.
Requires python-docx and Pillow. Run from repository root after evidence capture.
The archived input is immutable so regeneration does not duplicate additions.
"""
from pathlib import Path
from copy import deepcopy
import json, re, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'DOC/Nhom08_BaoCao_HQTCSDL_HomeFix.docx'
ARCHIVE = ROOT / 'DOC/archive/Nhom08_BaoCao_TruocBoSung_20261016.docx'
ARCHIVE.parent.mkdir(exist_ok=True)
if not ARCHIVE.exists(): shutil.copy2(OUT, ARCHIVE)
d = Document(ARCHIVE)
original_paragraphs = list(d.paragraphs)
original_tables = list(d.tables)
meta = json.loads((ROOT/'DOC/evidence/schema.json').read_text(encoding='utf-8'))
evidence = ROOT/'DOC/evidence/20261016-report'
screens = evidence/'screenshots'

def replace(p, text):
    # Preserve paragraph properties and the first run's formatting.
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if props is not None: r._r.insert(0, props)

def para(anchor, text, style='HF Body'):
    p=d.add_paragraph(text, style=style)
    anchor.addprevious(p._p)
    return p

def heading(anchor, text, level=2): return para(anchor,text,f'Heading {level}')

def celltext(cell,text):
    p=cell.paragraphs[0]
    replace(p,str(text))
    for extra in list(cell.paragraphs)[1:]: extra._p.getparent().remove(extra._p)

def table(anchor, caption, headers, rows):
    if caption: para(anchor,caption,'HF Table Caption')
    t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
    width=16/len(headers)
    for col in t.columns: col.width=Cm(width)
    props=deepcopy(original_tables[95]._tbl.tblPr)
    t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,props)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for c,text in zip(cells,row):
            c.width=Cm(width); c.text=str(text)
            p=c.paragraphs[0]; p.paragraph_format.line_spacing=1.1
            p.paragraph_format.space_after=Pt(3)
            for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10.5); r.bold=i==0
            if i==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EFF5'); c._tc.get_or_add_tcPr().append(sh)
        if i==0:
            repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    anchor.addprevious(t._tbl)
    return t

def setcode(t, text):
    cell=t.cell(1,0)
    for p in list(cell.paragraphs)[1:]: p._p.getparent().remove(p._p)
    for i,line in enumerate(text.strip().splitlines()):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.clear(); p.paragraph_format.line_spacing=Pt(11.5); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9.5)
        r.font.color.rgb=RGBColor.from_string('008000' if line.lstrip().startswith('--') else '000000')

def code(anchor,title,text):
    el=deepcopy(original_tables[43]._tbl); anchor.addprevious(el)
    from docx.table import Table
    t=Table(el,d._body); celltext(t.cell(0,0),title); setcode(t,text)

def picture(anchor,path,caption,max_height=10):
    path=Path(path)
    with Image.open(path) as im: w,h=im.size
    width=min(16,max_height*w/h)
    p=para(anchor,'','Normal'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    para(anchor,caption,'HF Figure Caption')

def anchor(text): return next(p._p for p in original_paragraphs if p.text==text)

# Fix facts while preserving the original structure and paragraph styling.
updates={
146:'Hệ thống tiếp nhận yêu cầu sửa chữa, lưu thông tin liên hệ, địa điểm và thời gian hẹn. Điều phối viên lập báo giá sơ bộ; chỉ sau khi khách đồng ý, đơn mới chuyển sang chờ phân công. Kỹ thuật viên nhận lệnh, cập nhật tiến độ, đề xuất vật tư và lập phiếu nghiệm thu.',
154:'HomeFix phân định trách nhiệm giữa khách hàng, kỹ thuật viên và các bộ phận vận hành. Điều phối viên quản lý báo giá và phân công; kế toán xác nhận tiền và đối soát; CSKH tiếp nhận vấn đề sau dịch vụ; quản trị viên quản lý tài khoản, danh mục và cấu hình. Giám đốc theo dõi báo cáo và quyết định các đề xuất chính sách.',
324:'Tên CSDL mặc định trong cấu hình mẫu là HomeFix_DBMS_Nhom08; máy kiểm thử hiện dùng HomeFix_DBMS_Nhom08_TiengViet. Giá trị DB_NAME trong .env mới là tên được ứng dụng kết nối. Bộ SQL độc lập mặc định tạo HomeFix_DBMS_Nhom08_Import và từ chối tên đã tồn tại. Không chạy bộ cài mới lên CSDL đang dùng; khi nâng cấp, sao lưu trước rồi chạy npm run db:init.',
560:'Lượt kiểm thử ngày 16/10/2026 ghi nhận 108 ca API/SQL đạt (36 + 72), không thất bại hoặc bỏ qua; 92 lượt kiểm tra màn hình ở hai kích thước desktop/mobile đạt và kịch bản nghiệp vụ gồm 17 thao tác đạt. Hai bộ cài SQL nguyên khối và module đều được cài trên CSDL mới, chạy lại migration và kiểm tra bảo toàn dữ liệu. Phục hồi Full–Diff–Log sang CSDL riêng đạt trạng thái ONLINE và vượt qua DBCC CHECKDB. Đây là kết quả trong phạm vi bộ kiểm thử, không phải cam kết hệ thống không còn lỗi.',
564:'Bằng chứng của lượt chạy nằm trong DOC/evidence/20261016-report. Cấu trúc hiện tại gồm 30 bảng, 7 thủ tục nghiệp vụ, 5 hàm, 6 view, 8 trigger, 33 chỉ mục ngoài khóa chính và 7 database role. Năm thủ tục ghi nghiệp vụ có transaction và xử lý lỗi. sp_SaoLuuHeThong được cài riêng khi chạy script backup, không tính vào bảy thủ tục nghiệp vụ.',
577:'Môi trường kiểm chứng là máy Windows cục bộ, SQL Server và Edge; giao diện được kiểm tra ở viewport desktop và mobile giả lập, chưa đại diện cho mọi điện thoại thật. Chưa đo tải nhiều người dùng hoặc thời gian phục hồi mục tiêu trong sự cố vận hành. Chuỗi Full–Diff–Log đã được phục hồi thực tế; STOPAT mới là ví dụ hướng dẫn, chưa được kiểm thử tự động. Khóa ứng dụng chung có thể hạn chế khả năng xử lý song song khi quy mô tăng.',
578:'Chuyển khoản vẫn cần kế toán đối chiếu và xác nhận; vị trí cập nhật theo thao tác người dùng. OTP/email dùng nhà cung cấp giả lập trong kiểm thử; AI bên ngoài không được gọi. Quyết định đề xuất chính sách được lưu nhưng chưa tự triển khai lịch hiệu lực hoặc cập nhật giá. Tài khoản kết nối nghiệp vụ chưa được tách hoàn toàn theo quyền tối thiểu của từng phân hệ; lịch sao lưu và sao lưu ảnh cần được thiết lập riêng khi vận hành.',
580:'Ưu tiên tiếp theo là giới hạn quyền tài khoản dịch vụ, lập lịch sao lưu và diễn tập phục hồi đồng bộ CSDL–ảnh, xác định RPO/RTO bằng phép đo, kiểm thử tải và giám sát lỗi. Có thể thu hẹp khóa theo đơn hoặc ví để tăng đồng thời, đồng thời giữ kiểm tra trạng thái và phiên bản tại SQL Server.',
591:'[9] Nhóm 08. Mã nguồn HomeFix, script SQL và kết quả kiểm thử. https://github.com/bndt0103/DBMS_HomeFix. Phiên bản tích hợp a79dff9 và các chỉnh sửa cục bộ; đối chiếu ngày 16/10/2026.',
584:'[2] Nguyễn Thị Phương Linh. Quản lý cửa hàng kinh doanh đồ trang trí nội thất Phúc Chí Thành. Đồ án tốt nghiệp, Trường Đại học Dân lập Hải Phòng, 2009. Tham khảo cách gắn đặc tả giao diện với tác nhân, thao tác và kết quả; không sử dụng công nghệ cũ làm căn cứ triển khai. https://lib.hpu.edu.vn/bitstream/handle/123456789/18404/87_NguyenThiPhuongLinh_CT902.pdf',
}
for i,text in updates.items(): replace(original_paragraphs[i],text)

# Refresh SQL listings to current metadata, including the changed wallet procedure.
modules={m['name']:m['definition'] for m in meta['modules']}
for t in original_tables:
    if len(t.columns)==1 and len(t.rows)==2 and t.cell(0,0).text in modules:
        setcode(t,modules[t.cell(0,0).text])

# Add the previously absent signature column with the same row formatting.
t=original_tables[17]; row=deepcopy(t.rows[-1]._tr); t._tbl.append(row)
for c,txt in zip(t.rows[-1].cells,['ChuKyKTVId','int','FK → TepDinhKem.Id; cho phép NULL','Tệp chữ ký của kỹ thuật viên lập nghiệm thu']): celltext(c,txt)
for r in t.rows:
    if r.cells[0].text=='ChuKyId': celltext(r.cells[3],'Tệp chữ ký khách hàng khi xác nhận nghiệm thu')

# Front matter: retain the cover and existing section/page numbering.
a=anchor('MỤC LỤC')
p=para(a,'LỜI MỞ ĐẦU','HF Front Title'); p.paragraph_format.page_break_before=True
for text in [
    'Một yêu cầu sửa chữa tại nhà không kết thúc ở việc phân công thợ. Báo giá, vật tư, nghiệm thu và các khoản tiền phải được ghi nhận thống nhất để khách hàng biết mình đã đồng ý điều gì, kỹ thuật viên xác định công việc cần làm và bộ phận vận hành có cơ sở đối chiếu. Khi mỗi bộ phận lưu thông tin riêng, việc truy lại một quyết định hoặc một khoản thu dễ trở nên thiếu căn cứ.',
    'Từ bài toán đó, nhóm xây dựng HomeFix trong học phần Hệ quản trị cơ sở dữ liệu. Trọng tâm của đồ án là thiết kế dữ liệu và hiện thực các cơ chế bảo vệ dữ liệu bằng SQL Server; ứng dụng web là môi trường để thực thi và kiểm chứng các cơ chế ấy. Báo cáo đi từ yêu cầu nghiệp vụ, mô hình dữ liệu đến chương trình SQL, giao dịch, phân quyền, sao lưu và kiểm thử.',
    'Nhóm cảm ơn ThS. Phan Thị Thể đã hướng dẫn học phần và cung cấp các yêu cầu thực hiện đồ án. Các kết quả được trình bày kèm mã nguồn, dữ liệu kiểm thử và ảnh chụp hệ thống. Những nội dung chưa triển khai hoặc chưa kiểm chứng được nêu riêng để phân biệt với chức năng đã hoàn thành.',
]: para(a,text)
para(a,'Nhóm 08 — Thành phố Hồ Chí Minh, tháng 10 năm 2026').alignment=WD_ALIGN_PARAGRAPH.RIGHT

a=anchor('3.2. Ràng buộc toàn vẹn')
heading(a,'3.1.1. Bộ SQL theo chức năng và quản lý phiên bản',3)
para(a,'SRC/database là nguồn chuẩn của cấu trúc và migration. Lệnh npm run sql:build sinh đồng thời bộ nguyên khối và các module dưới đây; npm run sql:check phát hiện bản sinh chưa đồng bộ. Chạy CSDL_HomeFix_Modules.sql để bảo đảm thứ tự phụ thuộc, không chọn chạy rời ngẫu nhiên từng file. Bộ backup/recovery được chạy riêng vì có thao tác quản trị CSDL.')
table(a,'Bảng 3.4. Bộ SQL được chia theo chức năng',['File trong SQL/modules','Nội dung'],[
['00_schema / 01_index','Bảng, khóa, ràng buộc / 20 lệnh CREATE INDEX tường minh'],['02_function / 03_view','5 hàm / 6 khung nhìn'],['04_sp / 05_transaction','2 thủ tục đọc / 5 thủ tục ghi có transaction'],['06_trigger / 07_security','8 trigger / 7 role, GRANT, REVOKE, DENY'],['08_seed / 09_data_migration','Dữ liệu mẫu / chuyển đổi dữ liệu bổ trợ'],['10_backup_recovery','Sao lưu và phục hồi; không chạy tự động khi cài mới']])
para(a,'33 index ngoài PK bao gồm cả index duy nhất do ràng buộc UNIQUE tạo ra; không đồng nhất con số này với 20 câu CREATE INDEX trong file module. Kiểm thử hai cách cài đối chiếu định nghĩa của 26 đối tượng lập trình SQL, cấu trúc cột, index và dữ liệu mẫu.')

a=anchor('3.4. Khung nhìn phục vụ truy vấn')
heading(a,'3.3.8. Thông báo đề xuất chính sách — trg_DeXuatChinhSach_ThongBao',3)
trigger=next((m for m in modules if m.startswith('trg_') and 'ChinhSach' in m),None)
if trigger:
    # Use the actual object name rather than assuming a migration's historical name.
    p=next(p for p in d.paragraphs if p.text.startswith('3.3.8.')); replace(p,'3.3.8. Thông báo đề xuất chính sách — '+trigger)
    para(a,'Trigger tạo thông báo cho người liên quan khi đề xuất chính sách phát sinh hoặc thay đổi trạng thái. Cách xử lý theo inserted/deleted giúp phản ánh quyết định trên dữ liệu thật, thay vì hiển thị danh sách đề xuất cố định trên dashboard. Quyền quyết định, lý do và phiên bản được API kiểm tra trước khi ghi; thông báo không tự thay đổi giá dịch vụ.')
    code(a,trigger,modules[trigger])

a=anchor('4.5. Xác thực và bảo vệ dữ liệu đầu vào')
heading(a,'4.4.1. Ma trận quyền và giới hạn của ngữ cảnh thực thi',3)
table(a,'Bảng 4.3. Quyền SQL tiêu biểu theo bộ phận',['Role','Quyền được cấp','Giới hạn'],[
['KH','Đọc dịch vụ công khai; EXEC sp_TaoDonHang','API giới hạn khách hàng của đơn'],['KTV','Đọc dịch vụ công khai','Nghiệp vụ nhận đơn/tiến độ qua API'],['DPV','Đọc tổng hợp đơn, hiệu suất; báo cáo tổng hợp','Không duyệt đối soát'],['CSKH','Đọc hỗ trợ/đơn; EXEC sp_XuLyHoTro và báo cáo','Không duyệt ví'],['KT','Đọc ví/doanh thu; duyệt ví, đối soát, báo cáo','Không sửa/xóa sổ ví'],['GD','Đọc doanh thu/hiệu suất; báo cáo tổng hợp','Quyền SELECT nhật ký đã REVOKE'],['ADMIN','Đọc view/nhật ký; CRUD danh mục; báo cáo','Vẫn DENY đọc cột mật khẩu và sửa/xóa sổ ví']])
para(a,'Các role này gắn với user WITHOUT LOGIN để dùng EXECUTE AS và kiểm chứng quyền [11]. Tài khoản đăng nhập web nằm trong NguoiDung, không phải login SQL Server. Chỉ luồng báo cáo hiện chuyển ngữ cảnh SQL theo vai trò; các nghiệp vụ khác dùng kết nối cấu hình và kiểm tra quyền tại API. Vì vậy, không xem role SQL minh họa là cơ chế phân quyền theo từng dòng cho toàn bộ ứng dụng.')
heading(a,'4.4.2. Minh chứng GRANT, REVOKE và DENY',3)
para(a,'Kiểm thử SQL xác nhận role báo cáo được đọc dữ liệu tổng hợp, không được đọc mật khẩu hoặc thực hiện nghiệp vụ tài chính của bộ phận khác. REVOKE loại bỏ quyền đã cấp trực tiếp; nó không tương đương DENY và không tự loại bỏ quyền được thừa hưởng từ một role khác.')
code(a,'Kiểm chứng quyền trong ngữ cảnh giám đốc',"EXECUTE AS USER = 'hf_GD';\nSELECT HAS_PERMS_BY_NAME('dbo.sp_BaoCaoTongHop',\n    'OBJECT', 'EXECUTE') AS DuocBaoCao; -- 1\nSELECT HAS_PERMS_BY_NAME('dbo.NhatKy',\n    'OBJECT', 'SELECT') AS DuocDocNhatKy; -- 0\nSELECT HAS_PERMS_BY_NAME('dbo.GiaoDichVi',\n    'OBJECT', 'UPDATE') AS DuocSuaSoVi; -- 0\nREVERT;")
para(a,'Khi thực hành lệnh bị từ chối, đặt EXECUTE AS/REVERT trong TRY…CATCH và luôn REVERT ở nhánh lỗi để không để lại ngữ cảnh của người khác trên kết nối dùng chung. Không kết luận bảo mật chỉ dựa vào việc nút thao tác bị ẩn trên giao diện.')

a=anchor('CHƯƠNG 5. XÂY DỰNG ỨNG DỤNG VÀ KIỂM THỬ')
heading(a,'4.7. Sao lưu và phục hồi dữ liệu')
heading(a,'4.7.1. Phạm vi dữ liệu cần bảo vệ',3)
para(a,'Bộ SQL tái tạo cấu trúc và dữ liệu mẫu giúp cài lại bản demo, nhưng không chứa các chứng từ mới phát sinh. Bản backup SQL lưu trạng thái CSDL; ảnh sửa chữa, chứng từ và chữ ký được lưu tại SRC/backend/uploads nên cần sao lưu riêng. Một bản phục hồi chỉ có metadata TepDinhKem mà thiếu tệp ảnh chưa phải bản phục hồi đầy đủ của ứng dụng.')
table(a,'Bảng 4.4. Thành phần sao lưu và mục đích',['Thành phần','Cách bảo vệ','Kiểm tra sau phục hồi'],[
['CSDL SQL Server','Full, Differential, Transaction Log có CHECKSUM','ONLINE; CHECKDB; kiểm tra đơn, thanh toán, số dư'],['Thư mục uploads','Sao chép nhất quán với mốc CSDL; lưu ở vị trí khác máy','Ảnh, chữ ký, chứng từ mở được qua API'],['Cấu hình .env','Lưu bản bảo mật, tách khỏi tài liệu và Git','Kết nối đúng CSDL; khóa và dịch vụ đúng môi trường'],['Mã nguồn/SQL','Quản lý phiên bản, ghi commit và bản migration','Bộ source tương thích cấu trúc được phục hồi']])
heading(a,'4.7.2. Chiến lược và thủ tục sao lưu',3)
para(a,'Script SQL/CSDL_HomeFix_SaoLuuPhucHoi.sql chuyển nguồn sang RECOVERY FULL, cài sp_SaoLuuHeThong rồi tạo Full, Diff và Log trong cùng lượt thực hành. Thủ tục nhận @LoaiSaoLuu, @DuongDanThuMuc và trả @TepSaoLuu OUTPUT. Loại không thuộc FULL/DIFF/LOG bị từ chối; tên tệp chứa GUID để mỗi lần gọi tạo tệp riêng. Dịch vụ SQL Server cần quyền ghi thư mục trên máy chủ SQL.')
code(a,'Gọi thủ tục sau khi đã cài script backup',"DECLARE @Tep nvarchar(1000);\nEXEC dbo.sp_SaoLuuHeThong\n    @LoaiSaoLuu='FULL', @TepSaoLuu=@Tep OUTPUT;\nRESTORE VERIFYONLY FROM DISK=@Tep WITH CHECKSUM;\nSELECT @Tep AS TepSaoLuu;")
para(a,'Đề xuất khi vận hành: Full mỗi tuần, Diff cuối ngày, Log mỗi 15 phút, lưu bản sao ngoài máy và định kỳ thử phục hồi. Đây là lịch đề xuất, chưa phải tác vụ đã được cài. SQL Server Express không có SQL Server Agent; có thể lập lịch bằng Windows Task Scheduler gọi sqlcmd. RPO mục tiêu 15 phút chỉ có ý nghĩa khi các lượt Log thành công và được giám sát; RTO phải được đo bằng diễn tập, chưa có số cam kết trong đồ án.')
heading(a,'4.7.3. Trình tự phục hồi và kiểm tra',3)
para(a,'Trong SQLCMD Mode, sửa TenCSDL thành tên CSDL nguồn, đặt ThucHienPhucHoi=1 rồi chạy script. Script từ chối nếu đích <nguồn>_Restore đã tồn tại. Cấu hình tự động áp dụng cho bộ cài có một file dữ liệu và một file log; bố cục nhiều file cần thiết lập MOVE thủ công. Không dùng REPLACE để ghi đè dữ liệu đang sử dụng.')
table(a,'Bảng 4.5. Chuỗi phục hồi đã thực hiện',['Bước','Thao tác','Ý nghĩa'],[
['1','RESTORE Full WITH MOVE, NORECOVERY','Tạo file đích riêng; giữ trạng thái RESTORING'],['2','RESTORE Diff WITH NORECOVERY','Áp dụng bản vi sai thuộc đúng Full nền'],['3','RESTORE LOG WITH RECOVERY','Áp dụng log trong chuỗi; hoàn tất recovery'],['4','DBCC CHECKDB; kiểm tra dữ liệu','Kiểm tra tính toàn vẹn và đối chiếu dữ liệu đã backup']])
para(a,'NORECOVERY giữ CSDL sẵn sàng nhận bản tiếp theo; RECOVERY kết thúc chuỗi và mở CSDL [10]. Trong vận hành, cần toàn bộ các bản Log liên tục đến mốc phục hồi. Ví dụ STOPAT trong script minh họa phục hồi đến thời điểm; cần bản Log chứa mốc đó và cân nhắc tail-log nếu nguồn còn truy cập được. Bài kiểm thử hiện xác minh chuỗi Full–Diff–Log, chưa xác minh STOPAT.')
heading(a,'4.7.4. Kết quả thực nghiệm và quy trình vận hành',3)
para(a,'tests/sql-bundle.test.js tạo CSDL tạm riêng, chạy bộ cài và chuỗi backup/restore, đối chiếu dữ liệu phục hồi, xác nhận trạng thái ONLINE và CHECKDB thành công. RESTORE VERIFYONLY được dùng để kiểm tra bản sao nhưng không thay thế thử phục hồi. Bằng chứng lưu tại sql-bundle.json và backup-recovery.txt trong thư mục bằng chứng của lượt chạy.')
para(a,'Khi khôi phục hệ thống: dừng ghi nghiệp vụ, giữ bản hiện trạng nếu còn đọc được, phục hồi sang CSDL mới, khôi phục uploads đúng mốc, đối chiếu ví–thanh toán–đơn và ảnh, sau đó đổi DB_NAME rồi chạy lại API. Chỉ chuyển người dùng sang bản phục hồi sau khi các kiểm tra đạt. Trước nâng cấp CSDL cục bộ, nhóm đã tạo bản COPY_ONLY có CHECKSUM và VERIFYONLY; tài khoản và dấu kiểm tra mật khẩu được bảo toàn sau migration.')

# Update existing evidence numbers and feature screenshots.
usage=json.loads((evidence/'sql-usage.json').read_text(encoding='utf-8'))
counts={x['procedureName']:x['executionCount'] for x in usage['procedures']}
for row in original_tables[96].rows[1:]:
    if row.cells[0].text in counts: celltext(row.cells[1],counts[row.cells[0].text])
for t in original_tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for old,new in [('87 kiểm thử','108 kiểm thử'),('26 trang','92 lượt màn hình'),('16 thao tác','17 thao tác'),('7 trigger','8 trigger')]:
                    if old in p.text: replace(p,p.text.replace(old,new))

image_replacements={'Hình 5.2.':'01-login-desktop.png','Hình 5.3.':'kh-home-desktop.png','Hình 5.4.':'dpv-home-desktop.png','Hình 5.5.':'ktv-home-desktop.png','Hình 5.6.':'kt-finance-desktop.png','Hình 5.7.':'cskh-home-desktop.png','Hình 5.8.':'admin-admin-users-desktop.png','Hình 5.9.':'gd-home-desktop.png'}
for i,p in enumerate(original_paragraphs):
    if p.style.name!='HF Figure Caption': continue
    name=next((v for k,v in image_replacements.items() if p.text.startswith(k)),None)
    if name and (screens/name).exists():
        imgp=original_paragraphs[i-1]
        if imgp._p.xpath('.//w:drawing'):
            imgp.clear(); imgp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            imgp.add_run().add_picture(str(screens/name),width=Cm(16))

a=anchor('CHƯƠNG 6. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN')
heading(a,'5.6.1. Môi trường, cấu hình và quy trình cài đặt',3)
table(a,'Bảng 5.6. Môi trường cài đặt bản DBMS',['Thành phần','Yêu cầu / thực tế kiểm chứng'],[
['Máy chủ','Windows x64; khuyến nghị RAM 8 GB trở lên cho ứng dụng và SQL cục bộ, không phải ngưỡng đã benchmark'],['Runtime','Node.js ≥ 22.12 và npm; lượt chạy dùng Node.js 24'],['CSDL','SQL Server Database Engine; máy kiểm thử localhost\\SQLEXPRESS, phiên bản 17'],['Kết nối','ODBC Driver 17 x64 với DB_AUTH=windows; SSMS hoặc sqlcmd để quản trị'],['Trình duyệt','Edge; Playwright dùng viewport 1440×1000 và 390×844'],['Dịch vụ tùy chọn','Gmail/Resend cho email OTP; Gemini cho chat nếu cấu hình; không bắt buộc để đăng nhập demo']])
para(a,'Bước 1. Cài Node.js, SQL Server Database Engine và ODBC phù hợp. Kiểm tra node --version, npm.cmd --version và dịch vụ SQL Server (SQLEXPRESS). SSMS chỉ là công cụ quản trị, không thay thế Database Engine. Giải nén mã nguồn vào thư mục có quyền ghi và mở terminal tại thư mục gốc.')
para(a,'Bước 2. Với cấu hình mặc định, chạy CAI_DAT.bat. Script cài thư viện, tạo .env nếu thiếu, sinh khóa JWT ngẫu nhiên, khởi tạo CSDL và build frontend. Nếu dùng instance/tên CSDL khác, chỉnh SRC/backend/.env trước bước db:init; không ghi mật khẩu SQL hoặc khóa email vào báo cáo.')
code(a,'Cài thủ công trong thư mục SRC',"npm.cmd install\n# Tạo backend/.env theo backend/.env.example; đặt JWT_SECRET riêng.\nnpm.cmd run db:init\nnpm.cmd run build\nnpm.cmd start")
para(a,'Bước 3. Mở http://localhost:3000, kiểm tra GET /api/health trả status=ok và đăng nhập tài khoản mẫu. Dừng bằng Ctrl+C. Các lần dùng sau chạy CHAY_HOMEFIX_DBMS.bat. Nếu frontend thay đổi, build lại trước khi start. Không có APK hoặc bộ cài desktop đóng gói trong phiên bản DBMS này.')
table(a,'Bảng 5.7. Xử lý lỗi cài đặt thường gặp',['Hiện tượng','Kiểm tra và cách xử lý'],[
['Không kết nối SQL','Dịch vụ SQL đang chạy; DB_SERVER đúng instance; tài khoản Windows có quyền'],['Không thấy ODBC Driver 17','Cài driver x64 hoặc cấu hình DB_ODBC_DRIVER theo driver đã cài'],['Cổng 3000 đang được dùng','Dừng phiên HomeFix cũ hoặc đổi PORT; dùng URL và CLIENT_ORIGINS tương ứng'],['Chạy SQL báo lỗi :setvar/:r','Bật SQLCMD Mode; chỉnh ThuMucSQL đúng thư mục SQL'],['Đăng nhập được nhưng OTP không gửi','Cấu hình provider và thông tin xác thực email; xem lỗi, không coi là đã gửi'],['Ảnh không mở sau phục hồi','Kiểm tra thư mục uploads đúng bản sao và quyền đọc của tiến trình']])
heading(a,'5.6.2. Kiểm tra sau cài đặt và tài liệu trợ giúp',3)
para(a,'Đăng nhập lần lượt bảy vai trò; kiểm tra danh mục, trạng thái đơn, ví và các trang báo cáo. CSDL mới có dữ liệu mẫu nên báo cáo đơn/doanh thu có thể rỗng; không tự điền số liệu để làm đẹp màn hình. Tệp HELP/HomeFix_HuongDan.chm cung cấp hướng dẫn cài đặt và thao tác từng bộ phận. Mục lục bên trái dùng tiếng Việt không dấu; nội dung dùng tiếng Việt có dấu và co giãn theo chiều rộng khung đọc.')

heading(a,'5.7. Ca kiểm thử nghiệp vụ có đầu vào và kết quả')
workflow=json.loads((evidence/'ui-workflow-results.json').read_text(encoding='utf-8'))
order='HF-'+str(workflow['orderId']).zfill(6)
para(a,f'Các ca sau được thực hiện trên CSDL tách biệt, sử dụng đơn {order} của lượt chạy. Tài khoản là dữ liệu demo, ảnh sửa chữa/chữ ký tải lên là tệp kiểm thử để xác minh lưu và phân quyền, không phải hồ sơ khách hàng thật. Cấu trúc trình bày tham khảo mẫu cài đặt–kiểm thử do nhóm cung cấp và cách đặc tả tương tác trong tài liệu [2]; nhãn nút, dữ liệu và kết quả dưới đây lấy từ bản DBMS hiện tại.')
cases=[
('Đặt dịch vụ','KH đã đăng nhập; dịch vụ Id=1 đang hoạt động.','Địa chỉ: 1 Võ Văn Ngân, Thủ Đức, TP.HCM. Mô tả: [UI TEST] Máy lạnh không làm mát. Đề nghị kiểm tra tại nhà.','Mở dịch vụ, điền địa chỉ và mô tả; bấm Gửi yêu cầu đặt dịch vụ.','Tạo đơn mới, trạng thái Chờ tiếp nhận; điều phối truy cập đúng đơn.','Đặt dịch vụ — dữ liệu nhập','Đặt dịch vụ — kết quả'),
('Báo giá sơ bộ và khách phê duyệt','DPV và KH truy cập cùng đơn đang chờ tiếp nhận.','Chẩn đoán: Kiểm tra máy lạnh không mát; báo giá công kiểm tra và sửa chữa. Phí lấy từ danh mục hiện tại.','DPV chọn Lập báo giá sơ bộ, gửi xác nhận. KH tải lại đơn, đọc báo giá và chọn Đồng ý báo giá.','Đơn chuyển Chờ duyệt báo giá rồi Chờ phân công; chưa có lệnh phân công trước khi khách đồng ý.','Lập báo giá sơ bộ — dữ liệu trước khi gửi','Đồng ý báo giá — kết quả'),
('Phân công và nhận lệnh','Đơn chờ phân công; KTV Minh sẵn sàng và đủ điều kiện được hệ thống gợi ý.','Chọn Kỹ thuật viên Minh trong danh sách hợp lệ.','DPV chọn Phân công kỹ thuật viên và xác nhận. KTV mở đơn, chọn Chấp nhận đơn rồi Xác nhận nhận đơn.','Đơn chuyển Chờ thợ nhận rồi Thợ đã nhận; người được phân công nhìn thấy thông tin đúng đơn.','Phân công kỹ thuật viên — dữ liệu trước khi gửi','Chấp nhận đơn — kết quả'),
('Tiến độ và phê duyệt vật tư','KTV đã nhận đơn; thực hiện lần lượt di chuyển, đến nơi, bắt đầu xử lý.','Vật tư: Van máy lạnh; số lượng mặc định 1; đơn giá 220.000 đồng; bảo hành 6 tháng; ghi chú Van cũ bị hỏng, cần thay mới.','KTV Đề xuất vật tư, Gửi khách hàng duyệt. KH tải lại đơn và chọn Đồng ý thay vật tư.','Phiếu vật tư được khách duyệt; giá trị 220.000 đồng là căn cứ lập nghiệm thu, chưa phải khoản tiền đã thu.','Đề xuất vật tư — dữ liệu trước khi gửi','Đồng ý thay vật tư — kết quả'),
('Lập phiếu nghiệm thu và hai chữ ký','Đơn đang xử lý; vật tư được phê duyệt.','Nguyên nhân: Van máy lạnh đã hỏng gây giảm khả năng làm mát. Biện pháp: Đã thay van, kiểm tra kín và chạy thử ổn định. Ảnh và chữ ký kiểm thử.','KTV lập phiếu, vẽ chữ ký, xem lại rồi gửi. KH đọc phiếu, xác nhận đã kiểm tra và tải chữ ký khi Xác nhận nghiệm thu.','Trước khách ký chỉ có chữ ký KTV; sau xác nhận hiển thị hai chữ ký riêng. Đơn hoàn thành, trạng thái thanh toán vẫn phụ thuộc chứng từ thu tiền.','Hoàn thành · Lập phiếu nghiệm thu — dữ liệu trước khi gửi','Xác nhận nghiệm thu — kết quả'),
('Thu COD, đối soát và đánh giá','KH đã nghiệm thu; KTV xác nhận thu tiền thực tế; kế toán đăng nhập.','Xác nhận COD theo số tiền hệ thống tính; nhận xét: Kỹ thuật viên giải thích rõ, sửa đúng yêu cầu.','KTV chọn Xác nhận đã thu COD. KT mở Tài chính → ví/đối soát, chọn đúng mã đơn và Xác nhận xử lý. KH chọn Đánh giá dịch vụ.','Một khoản thanh toán và đối soát được xác nhận; hoa hồng ghi sổ ví. Khách gửi được đánh giá sau hoàn thành và thanh toán.','Xác nhận đã thu COD — dữ liệu trước khi gửi','Kế toán — kết quả đối soát'),
]
captures={x['label']:x['filename'] for x in workflow.get('reportCaptures',[])}
for i,(title,pre,inp,steps,result,img1,img2) in enumerate(cases,1):
    heading(a,f'5.7.{i}. {title}',3)
    for label,text in [('Điều kiện',pre),('Input',inp),('Thực hiện',steps),('Output thực tế',result)]:
        p=para(a,label+': '+text)
    # One readable viewport per case; all before/after captures are supplied in evidence.
    label=img1 if i in (1,2,4,5) else img2
    if label in captures: picture(a,screens/captures[label],f'Hình 5.{9+i}. {title} — ảnh chụp trong lượt kiểm thử',10.5)
heading(a,'5.7.7. Ca không hợp lệ, quyền và dữ liệu đồng thời',3)
table(a,'Bảng 5.8. Các ca kiểm thử bổ sung và tiêu chí chấp nhận',['Ca / Input','Output cần đạt','Minh chứng'],[
['Thiếu dữ liệu bắt buộc hoặc số tiền âm','Từ chối dữ liệu; không tạo chứng từ sai','API/SQL tests'],['Khách truy cập đơn của người khác','Từ chối truy cập dữ liệu ngoài quyền sở hữu','api.test.js'],['Phân công trước khi duyệt báo giá hoặc thợ không hợp lệ','Không tạo lệnh phân công','api.test.js'],['Hai phiên SQL duyệt cùng yêu cầu ví','Một lần thành công; một bị từ chối; một bút toán','dbms.test.js'],['Lỗi chèn lịch sử sau khi tạo đơn','Rollback; không giữ đơn thiếu lịch sử','dbms.test.js'],['INSERT sổ ví nhiều dòng / làm âm số dư','Cộng đúng tổng / rollback toàn bộ khi sai','dbms.test.js'],['OTP sai, hết hạn, dùng lại hoặc khác mục đích','Từ chối; sai 5 lần khóa mã; đổi mật khẩu thu hồi phiên','otp.test.js'],['Cập nhật GPS rồi từ chối quyền vị trí','Bản đồ KTV đổi điểm đến; khi từ chối giữ tọa độ cũ','ui-workflow.mjs'],['Quyết định chính sách sai quyền/phiên bản cũ','Không ghi quyết định không hợp lệ','director-reports.test.js'],['Cài SQL vào tên CSDL đã tồn tại','Dừng, không ghi đè CSDL','sql-bundle.test.js']])
para(a,'Ảnh trước/sau của từng thao tác lưu cùng manifest reportCaptures trong ui-workflow-results.json. Smoke test kiểm tra mở màn hình, lỗi JavaScript/API và tràn ngang; không thay thế kiểm thử thao tác. Khi ghi kết quả báo cáo, chỉ đánh dấu đạt cho ca có assertion thành công trong lượt chạy, không suy ra kết quả từ ảnh giao diện đơn thuần.')

# Final chapter additions and references.
a=anchor('6.2. Hạn chế')
para(a,'Bản bàn giao bổ sung bộ SQL chia chức năng, cơ chế kiểm tra đồng bộ với nguồn, kiểm thử bộ cài và phục hồi thực tế. Dashboard sử dụng đề xuất chính sách và hoa hồng đã đối soát từ dữ liệu thật; chữ ký kỹ thuật viên và khách hàng được phân biệt. Hướng dẫn CHM và hồ sơ kiểm thử giúp người khác cài lại và lần theo từng nghiệp vụ để kiểm chứng.')
a=anchor('PHỤ LỤC A. PHÂN CÔNG THỰC HIỆN')
for text in [
'[10] Microsoft Learn. Restore and recovery overview (SQL Server). https://learn.microsoft.com/en-us/sql/relational-databases/backup-restore/restore-and-recovery-overview-sql-server. Truy cập 16/10/2026.',
'[11] Microsoft Learn. CREATE USER (Transact-SQL). https://learn.microsoft.com/en-us/sql/t-sql/statements/create-user-transact-sql. Truy cập 16/10/2026.',
'[12] Khoa CNTT, Trường Đại học Nông Lâm TP.HCM. Đề cương chi tiết môn Cơ sở dữ liệu nâng cao. Tham khảo định hướng trình bày transaction và điều khiển truy xuất đồng thời; không thay thế rubric DBMS330284. https://fit.hcmuaf.edu.vn/fit-3901-1/vn/de-cuong-chi-tiet-mon-co-so-du-lieu-nang-cao.html. Truy cập 16/10/2026.',
'[13] Nhóm 08. Mẫu cài đặt và kiểm thử hệ thống CNPM (Mau_Caidatkiemthu.docx). Tài liệu nội bộ do nhóm cung cấp; tham khảo hình thức trình bày, điều chỉnh theo mã nguồn DBMS.',
]: para(a,text)

# Append after existing appendix without changing section geometry.
a=OxmlElement('w:p'); d._element.body.insert(len(d._element.body)-1,a)
heading(a,'PHỤ LỤC C. ĐỐI CHIẾU RUBRIC VÀ HỒ SƠ KIỂM CHỨNG',1)
table(a,'Bảng C.1. Ánh xạ tiêu chí chấm điểm với sản phẩm',['Tiêu chí / trọng số','Vị trí báo cáo','Bằng chứng'],[
['Thiết kế CSDL / 15%','Chương 2; 30 bảng, quan hệ và chuẩn hóa','schema.json; sơ đồ ERD'],['Ràng buộc, trigger, view, index / 15%','Chương 3; phân tích và đo chỉ mục','SQL/modules; index-benchmark.json'],['SP, function, xử lý lỗi / 15%','Mục 3.5–3.6 và 4.2','26 định nghĩa SQL; sql-usage.json'],['Transaction, bảo mật / 10%','Chương 4; 5 giao dịch, 7 role','dbms.json; backup/restore'],['Ứng dụng SQL / 20%','Chương 5; API gọi SP/function','api-results.json; 17 bước UI'],['Giao diện / 5%','Mục 5.3 và 5.7','92 lượt desktop/mobile; ảnh'],['Báo cáo, tài liệu / 10%','Word/PDF, tham khảo, phụ lục','README; HELP/HomeFix_HuongDan.chm'],['Vấn đáp, làm việc nhóm / 10%','Phụ lục A; SQL/HUONG_DAN_VAN_DAP.md','Thành viên trình bày trực tiếp phần phụ trách']])
para(a,'Bảng này xác định nơi kiểm chứng từng tiêu chí, không tự quy đổi thành điểm. Năng lực trình bày và đóng góp cá nhân cần được nhóm xác nhận và thể hiện trong buổi bảo vệ. Số lượng đối tượng chỉ là điều kiện tối thiểu; chất lượng thiết kế, giải thích và chạy demo mới là căn cứ đánh giá.')
heading(a,'PHỤ LỤC D. HƯỚNG DẪN TRA CỨU VÀ TÁI LẬP',1)
table(a,'Bảng D.1. Các tệp dùng khi cài đặt và vấn đáp',['Tệp / lệnh','Cách dùng'],[
['HELP/HomeFix_HuongDan.chm','Mở trên Windows; chọn bộ phận ở mục lục bên trái'],['HELP/html/index.html','Bản HTML dự phòng cùng nội dung nếu máy không mở CHM'],['SQL/CSDL_HomeFix_Modules.sql','SQLCMD; tạo CSDL mới, thực thi module theo thứ tự'],['npm run db:init','Nâng cấp CSDL cấu hình; sao lưu trước khi áp dụng dữ liệu đang dùng'],['npm run test:sql-usage','Build, API/SQL, Extended Events, UI trên CSDL tách biệt'],['DOC/evidence/20261016-report','Kết quả JSON, ảnh và log backup; thời gian JSON theo UTC'],['SQL/HUONG_DAN_VAN_DAP.md','Bản đồ đối tượng và câu hỏi để đối chiếu mã nguồn']])
para(a,'Các tài khoản demo dùng đuôi @homefix.local: kh, ktv, dpv, cskh, kt, admin, gd; mật khẩu ban đầu HomeFix@123. Đây là dữ liệu mẫu để thực hành cục bộ, không dùng làm thông tin xác thực khi triển khai thật. Nếu tài khoản đã đổi mật khẩu, chạy lại khởi tạo không đặt lại mật khẩu cũ.')
para(a,'Quy trình kiểm tra hồ sơ: mở Word và PDF để đối chiếu mục lục, hình/bảng; mở CHM và thử các liên kết; chạy ứng dụng thật để kiểm tra kết nối SQL và nghiệp vụ. Khi đổi mã nguồn hoặc schema, sinh lại SQL và chạy kiểm thử trước khi cập nhật số liệu trong báo cáo.')
a.getparent().remove(a)

# Keep new captions sequential in reading order; Word rebuilds caption lists below.
for prefix,style in [('Hình','HF Figure Caption'),('Bảng','HF Table Caption')]:
    counters={}
    for p in d.paragraphs:
        if p.style.name!=style: continue
        m=re.match(rf'{prefix} ([0-9A-Z]+)\.\d+\.\s*(.*)',p.text)
        if not m: continue
        chapter,body=m.groups(); counters[chapter]=counters.get(chapter,0)+1
        replace(p,f'{prefix} {chapter}.{counters[chapter]}. {body}')

# Update fields on opening; Word COM is used to refresh actual values and pagination.
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); d.settings.element.append(update)
d.core_properties.title='HomeFix — Báo cáo đồ án Hệ quản trị cơ sở dữ liệu'
d.core_properties.subject='DBMS330284 — Nhóm 08; rà soát 16/10/2026'
d.save(OUT)
print(f'Saved {OUT}; paragraphs={len(d.paragraphs)}, tables={len(d.tables)}, images={len(d.inline_shapes)}')
